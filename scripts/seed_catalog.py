"""Carga única del catálogo de materiales desde un Excel (columnas Categoría,
Material, Densidad, Métrica; encabezado en la fila 4 del listado original).

Uso:
    python scripts/seed_catalog.py [ruta/al/listado_materiales.xlsx]

Escribe en la misma base que usa la app (Supabase si hay credenciales en
.env / secrets, si no el SQLite local). Es idempotente: si el material ya
existe (misma categoría y nombre) solo actualiza densidad y métrica."""
import sys
from pathlib import Path

import openpyxl
from dotenv import load_dotenv

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
load_dotenv(ROOT / ".env")

from db import repository as repo  # noqa: E402

DEFAULT_PATH = Path.home() / "Downloads" / "listado_materiales.xlsx"
# prefijos de los encabezados reales: "Densidad (kg/m³)", "Métrica que usa"...
HEADER_PREFIXES = {"categor": "category", "material": "name", "densidad": "density", "métrica": "metric", "metrica": "metric"}


def _classify(value) -> str | None:
    text = str(value or "").strip().lower()
    return next((field for prefix, field in HEADER_PREFIXES.items() if text.startswith(prefix)), None)


def read_rows(path: Path) -> list[tuple[str, str, float | None, str]]:
    sheet = openpyxl.load_workbook(path, data_only=True).active
    header_row, columns = None, {}
    for number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        found = {_classify(v): i for i, v in enumerate(row) if _classify(v)}
        if {"category", "name"} <= found.keys():
            header_row, columns = number, found
            break
    if not header_row:
        raise SystemExit("No encontré la fila de encabezados (Categoría / Material).")

    rows = []
    for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        category = str(row[columns["category"]] or "").strip()
        name = str(row[columns["name"]] or "").strip()
        if not category or not name:
            continue
        density = row[columns["density"]] if "density" in columns else None
        try:
            density = float(density) if density not in (None, "") else None
        except (TypeError, ValueError):
            density = None
        metric = str(row[columns["metric"]] or "").strip() if "metric" in columns else ""
        rows.append((category, name, density, metric))
    return rows


def main() -> None:
    path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_PATH
    if not path.exists():
        raise SystemExit(f"No existe el archivo: {path}")
    rows = read_rows(path)
    print(f"{len(rows)} materiales leídos de {path.name}. Base: {'SQLite local' if repo.USING_LOCAL else 'Supabase'}")
    for category, name, density, metric in rows:
        repo.upsert_material(category, name, density, metric)
    print(f"Listo: {len(rows)} materiales cargados/actualizados.")


if __name__ == "__main__":
    main()
